"""Tests for the excel-duplicate-handler merge helper."""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

REPO_ROOT = Path(__file__).parents[2]
SCRIPT = (
    REPO_ROOT
    / "plugins"
    / "report-ai-study-pipeline"
    / "skills"
    / "excel-duplicate-handler"
    / "scripts"
    / "merge_excel_duplicates.py"
)


def _write_workbook(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "_10_TST"
    worksheet.append(["FID", "TST_DAT1"])
    worksheet.append(["A", "2026-01-01"])
    worksheet["B2"].number_format = "DD/MM/YYYY"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _write_named_workbook(
    path: Path, sheet: str, headers: list[str], rows: list[list[str]]
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def test_batch_lock_temp_merge_uses_project_shaped_outputs(tmp_path: Path) -> None:
    dataset_dir = tmp_path / "data" / "raw" / "Mini" / "datasets"
    main = dataset_dir / "10_TST.xlsx"
    lock = dataset_dir / "~$10_TST.xlsx"
    artifact_root = tmp_path / "project"

    _write_workbook(main)
    before = main.read_bytes()
    lock.write_bytes(b"not a real workbook")

    result = subprocess.run(  # noqa: S603
        [
            sys.executable,
            str(SCRIPT),
            "--dataset-dir",
            str(dataset_dir),
            "--artifact-root",
            str(artifact_root),
        ],
        cwd=REPO_ROOT,
        check=True,
        capture_output=True,
        text=True,
    )

    merged = artifact_root / "data" / "raw" / "Mini" / "datasets" / "10_TST.xlsx"
    backup_dir = artifact_root / "data" / "raw" / "Mini" / "_dataset"
    backup = backup_dir / "10_TST.xlsx"
    backup_lock = backup_dir / "~$10_TST.xlsx"
    active_lock = artifact_root / "data" / "raw" / "Mini" / "datasets" / "~$10_TST.xlsx"
    report = artifact_root / "output" / "Mini" / "audit" / "datasets" / "10_TST" / "merge_report.md"
    provenance = (
        artifact_root / "output" / "Mini" / "audit" / "datasets" / "10_TST" / "merge_provenance.csv"
    )
    batch_report = artifact_root / "output" / "Mini" / "audit" / "dataset_duplicate_merge_report.md"

    assert "merged_groups=1" in result.stdout
    assert merged.read_bytes() == before
    assert backup.read_bytes() == before
    assert backup_lock.is_file()
    assert not active_lock.exists()
    assert main.read_bytes() == before
    assert report.is_file()
    assert provenance.is_file()
    assert batch_report.is_file()
    assert "dataset_workbook:" in report.read_text(encoding="utf-8")
    assert "raw_dataset_snapshot:" in report.read_text(encoding="utf-8")
    assert "invalid_source_count: 1" in report.read_text(encoding="utf-8")
    assert "removed_active_branch_file_count: 1" in report.read_text(encoding="utf-8")
    assert "`10_TST` | merged" in batch_report.read_text(encoding="utf-8")


def test_subset_branch_appends_into_superset_main_schema(tmp_path: Path) -> None:
    dataset_dir = tmp_path / "data" / "raw" / "Mini" / "datasets"
    main = dataset_dir / "14_Case_Control.xlsx"
    branch = dataset_dir / "14_CaseControl.xlsx"
    artifact_root = tmp_path / "project"

    _write_named_workbook(
        main,
        "_14_Case_Control",
        ["FID", "CC_BASE", "CC_RBSND"],
        [["M1", "main", "extra"]],
    )
    _write_named_workbook(
        branch,
        "_14_CaseControl",
        ["FID", "CC_BASE"],
        [["B1", "branch"]],
    )

    result = subprocess.run(  # noqa: S603
        [
            sys.executable,
            str(SCRIPT),
            "--study",
            "Mini",
            "--dataset",
            "14_Case_Control",
            "--main",
            str(main),
            "--branch",
            str(branch),
            "--artifact-root",
            str(artifact_root),
        ],
        cwd=REPO_ROOT,
        check=True,
        capture_output=True,
        text=True,
    )

    merged = artifact_root / "data" / "raw" / "Mini" / "datasets" / "14_Case_Control.xlsx"
    backup_dir = artifact_root / "data" / "raw" / "Mini" / "_dataset"
    backup = backup_dir / "14_Case_Control.xlsx"
    backup_branch = backup_dir / "14_CaseControl.xlsx"
    active_branch = artifact_root / "data" / "raw" / "Mini" / "datasets" / "14_CaseControl.xlsx"
    report = (
        artifact_root
        / "output"
        / "Mini"
        / "audit"
        / "datasets"
        / "14_Case_Control"
        / "merge_report.md"
    )

    assert "appended_rows=1" in result.stdout
    assert backup.is_file()
    assert backup_branch.is_file()
    assert not active_branch.exists()
    assert report.is_file()
    assert "appended_rows: 1" in report.read_text(encoding="utf-8")
    assert "removed_active_branch_file_count: 1" in report.read_text(encoding="utf-8")

    workbook = load_workbook(merged)
    try:
        worksheet = workbook["_14_Case_Control"]
        assert [worksheet.cell(row=1, column=idx).value for idx in range(1, 4)] == [
            "FID",
            "CC_BASE",
            "CC_RBSND",
        ]
        assert worksheet.max_row == 3
        assert worksheet.cell(row=3, column=1).value == "B1"
        assert worksheet.cell(row=3, column=2).value == "branch"
        assert worksheet.cell(row=3, column=3).value is None
    finally:
        workbook.close()


def test_partial_overlap_routes_to_human_review_without_outputs(tmp_path: Path) -> None:
    dataset_dir = tmp_path / "data" / "raw" / "Mini" / "datasets"
    main = dataset_dir / "18_2_TargConcom.xlsx"
    branch = dataset_dir / "18_1_TargConcom.xlsx"
    artifact_root = tmp_path / "project"

    _write_named_workbook(
        main,
        "_18_2_TargConcom",
        ["FID", "CONC_B", "ONLY_MAIN"],
        [["M1", "main", "main-only"]],
    )
    _write_named_workbook(
        branch,
        "_18_1_TargConcom",
        ["FID", "CONC_A", "ONLY_BRANCH"],
        [["B1", "branch", "branch-only"]],
    )

    result = subprocess.run(  # noqa: S603
        [
            sys.executable,
            str(SCRIPT),
            "--study",
            "Mini",
            "--dataset",
            "18_TargConcom",
            "--main",
            str(main),
            "--branch",
            str(branch),
            "--artifact-root",
            str(artifact_root),
        ],
        cwd=REPO_ROOT,
        check=True,
        capture_output=True,
        text=True,
    )

    merged = artifact_root / "data" / "raw" / "Mini" / "datasets" / "18_2_TargConcom.xlsx"
    active_branch = artifact_root / "data" / "raw" / "Mini" / "datasets" / "18_1_TargConcom.xlsx"
    backup = artifact_root / "data" / "raw" / "Mini" / "_dataset" / "18_2_TargConcom.xlsx"
    backup_branch = artifact_root / "data" / "raw" / "Mini" / "_dataset" / "18_1_TargConcom.xlsx"
    review = (
        artifact_root
        / "output"
        / "Mini"
        / "audit"
        / "human_review"
        / "excel"
        / "18_TargConcom"
        / "duplicate_review_report.md"
    )
    merge_report = (
        artifact_root
        / "output"
        / "Mini"
        / "audit"
        / "datasets"
        / "18_TargConcom"
        / "merge_report.md"
    )

    assert "status=human_review_required" in result.stdout
    assert "review_report=" in result.stdout
    assert merged.exists()
    assert active_branch.exists()
    assert backup.exists()
    assert backup_branch.exists()
    assert not merge_report.exists()
    assert review.is_file()
    review_text = review.read_text(encoding="utf-8")
    assert "not_handled_by_auto_merge" in review_text
    assert "no merged workbook was created" in review_text
    assert "partial_overlap" in review_text
